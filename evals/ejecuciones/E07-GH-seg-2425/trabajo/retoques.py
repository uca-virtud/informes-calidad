import re, sys, docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt

RUTA = "borrador.docx"
d = docx.Document(RUTA)

def pars():
    return [Paragraph(el, d._body) for el in d.element.body.iterchildren() if el.tag == qn("w:p")]

def find(prefix):
    for p in pars():
        if p.text.strip().startswith(prefix):
            return p
    sys.exit(f"Ancla no encontrada: {prefix}")

def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for h in p._p.findall(qn("w:hyperlink")):
        p._p.remove(h)

def nota_antes(anchor, text):
    el = OxmlElement("w:p")
    anchor._p.addprevious(el)
    p = Paragraph(el, anchor._parent)
    p.style = d.styles["Normal"]
    r = p.add_run(text)
    r.font.size = Pt(8)

# Frases modelo no acreditadas por las fuentes
p = find("Las modalidades de impartición del título se corresponden")
set_text(p, "Las modalidades de impartición del título se corresponden con la memoria verificada.")
p = find("En relación con el programa formativo del título de Grado en Historia")
set_text(p, "En el Grado en Historia, la coordinación docente se articula a través de las figuras de coordinador/a de asignatura y de coordinador/a del título.")
p = find("El desarrollo de la coordinación se asegura")
set_text(p, "Las reuniones de coordinación del curso 2024-25 y sus actas: [PENDIENTE: enlace a las actas de coordinación de 2024-25].")
p = find("Respecto a la gestión de los TFG/TFM")
for r in p.runs:
    r.text = r.text.replace("TFG/TFM", "TFG").replace("ejecución del TFM", "ejecución del TFG")
p = find("El PTGAS no se adscribe a títulos concretos")
set_text(p, "Como no hay PTGAS adscrito al título, el apoyo a la docencia lo prestan los servicios del campus y los servicios centrales de la UCA. La satisfacción del alumnado y del profesorado con los recursos (ISGC-P06-03 e ISGC-P06-04) se analiza en el subcriterio 5.1. La adecuación del personal de apoyo a las necesidades del título: [PENDIENTE: valoración del centro].")

# Restos de enlaces genéricos tras la depuración
for pref in ["ANEXO II:  1) Tabla Personal para impartir el título"]:
    for p in pars():
        if p.text.strip().startswith(pref):
            set_text(p, "ANEXO II: 1) Tabla Personal para impartir el título (último año impartido): véase el Anexo 2 de este autoinforme.")
p = find("https://data.uca.es/lincebi -")
set_text(p, "Fuente: Sistema de Información de la UCA (data.uca.es), informes ISGC-P05-01 y P05 (resto de indicadores).")

# Filas vacías en tablas de puntos débiles
for t in d.tables:
    if t.rows and t.rows[0].cells[0].text.strip() == "Puntos débiles":
        for row in list(t.rows)[1:]:
            if not "".join(c.text for c in row.cells).strip():
                row._tr.getparent().remove(row._tr)

# Mayúsculas en nombres de asignatura
MINUS = {"de", "del", "la", "las", "el", "los", "y", "e", "en", "a", "al", "por", "con", "o", "u"}
ROMANOS = {"i", "ii", "iii", "iv", "v", "vi"}
def capitaliza(nombre):
    pal = nombre.split(" ")
    out = []
    for k, w in enumerate(pal):
        base = w.lower().strip(":,")
        if base in ROMANOS:
            out.append(w.upper())
        elif k > 0 and base in MINUS:
            out.append(w.lower())
        else:
            out.append(w[:1].upper() + w[1:].lower())
    return " ".join(out)
for t in d.tables:
    if t.rows and t.rows[0].cells[0].text.strip() == "Asignatura":
        for row in list(t.rows)[1:]:
            c = row.cells[0]
            if c.text.startswith("GLOBAL"):
                continue
            for par in c.paragraphs:
                if par.runs:
                    txt = capitaliza(par.text)
                    par.runs[0].text = txt
                    for r in par.runs[1:]:
                        r._r.getparent().remove(r._r)

# Notas de fuente del Anexo 1
nota_antes(find("P01 – Difusión de la información."),
    "Fuente: Sistema de Información de la UCA (data.uca.es), extracción del 04/10/2026. Los valores de la columna 2020-21 de los "
    "indicadores P01 se han tomado del autoinforme anterior porque el Sistema de Información solo muestra los cuatro últimos cursos.")
nota_antes(find("3) P05 – Gestión del Personal Académico"),
    "Notas: las tasas de abandono, graduación y eficiencia son dinámicas y el Sistema de Información las recalcula cada año, por lo "
    "que pueden diferir de las publicadas en el autoinforme anterior. La serie ISGC-P04-13 de 2021-22 a 2023-24 es aproximadamente la "
    "mitad de la publicada en el autoinforme anterior; se reproduce tal como la ofrece el Sistema de Información. En 2020-21 la movilidad "
    "saliente internacional figura sin dato («-») y se consigna como 0%.")
nota_antes(find("4) Gestión de los recursos para el aprendizaje"),
    "Nota: en los datos del centro, la suma de ISGC-P05-06 e ISGC-P05-07 supera el 100% en algunos cursos; se reproducen tal como los ofrece el Sistema de Información.")
nota_antes(find("5) P07 – Resultados: A)"),
    "Nota: los valores 2020-21 de ISGC-P04-16 a ISGC-P04-19 y de los indicadores P06 proceden del autoinforme anterior.")
nota_antes(find("7) P07 – Resultados: C)"),
    "* No disponible: el Sistema de Información no ofrece dato de ISGC-P07-07 para la cohorte 2022-23. Los indicadores de egresados se "
    "refieren al año de egreso (cohortes 2019-20 a 2023-24) y se basan en tasas de respuesta bajas (en torno a 10 respuestas por cohorte).")
nota_antes(find("ANEXO 2:"),
    "Nota: el Sistema de Información muestra «-» cuando no hay registros en el BAU; se consigna como 0%.")

d.save(RUTA)
print("OK retoques")
