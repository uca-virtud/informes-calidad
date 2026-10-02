#!/usr/bin/env python3
"""Rellena la parte determinista del Autoinforme de Seguimiento (UCA/ACCUA).

Vuelca sobre la plantilla oficial .docx, sin alterar su formato:
  - Encabezado (curso, convocatoria, título, centro)
  - Tabla de datos de identificación del título
  - Tablas de indicadores del Anexo 1 (ISGC P01-P07), por TÍTULO y por CENTRO, con las
    cabeceras de años tomadas de `anios` y `anios_egresados`
  - (Opcional) Puntos fuertes/débiles por criterio y Plan de Mejora

Uso:
  python rellenar_autoinforme.py --plantilla plantilla.docx --datos datos.json \
         --salida autoinforme.docx

El esquema de datos.json está en references/mapa-indicadores-lincebi.md.
El análisis narrativo NO lo rellena este script: se redacta después sobre el .docx.
"""
import argparse
import json
import re
import sys

try:
    import docx
except ImportError:
    sys.exit("Falta python-docx. Instala con: pip install python-docx --break-system-packages")

YEAR_RE = re.compile(r"^\s*\d{4}-\d{2}\s*$")
CODE_RE = re.compile(r"ISGC-P\d{2}[-_]\d{2}(?:_\d{2})?")

# Mapa criterio -> (indice tabla puntos fuertes, indice tabla puntos debiles)
PUNTOS_TABLES = {
    1: (5, 6), 2: (8, 9), 3: (12, 13), 4: (15, 16),
    5: (18, 19), 6: (22, 23), 7: (24, 25),
}
PLAN_MEJORA_TABLE = 26


def set_cell_text(cell, text):
    """Reemplaza el contenido de una celda conservando el estilo del 1er párrafo.

    Soporta varias líneas (\n -> salto de línea) y elimina párrafos/placeholders
    sobrantes que trae la plantilla en la celda.
    """
    text = "" if text is None else str(text)
    lines = text.split("\n")
    # elimina párrafos adicionales de la celda (placeholders de la plantilla)
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    if p.runs:
        run = p.runs[0]
        run.text = lines[0]
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        run = p.add_run(lines[0])
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)


def distinct_leading(cells, first_year_col):
    """Textos distintos y no vacíos de las celdas anteriores a la 1ª columna-año."""
    out = []
    for c in cells[:first_year_col]:
        t = c.text.strip()
        if t and (not out or out[-1] != t):
            out.append(t)
    return out


def fill_header(doc, enc):
    if not enc:
        return
    for p in doc.paragraphs:
        t = p.text.strip()
        low = t.lower()
        if low.startswith("autoinforme de seguimiento curso") and enc.get("curso"):
            set_para(p, f"Autoinforme de seguimiento Curso {enc['curso']}")
        elif low.startswith("(convocatoria") and enc.get("convocatoria"):
            set_para(p, f"(Convocatoria {enc['convocatoria']})")
        elif t.upper().startswith("GRADO/MÁSTER EN") and enc.get("titulo"):
            set_para(p, f"{t} {enc['titulo']}")
        elif t.upper().startswith("CENTRO/S") and enc.get("centro"):
            set_para(p, f"{t} {enc['centro']}")


def set_para(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def fill_identity(doc, ident):
    if not ident:
        return
    # Busca la tabla de identificación (contiene "ID Ministerio")
    for tb in doc.tables:
        labels = [tb.cell(r, 0).text.strip() for r in range(len(tb.rows))]
        if any("ID Ministerio" in l for l in labels):
            for r, row in enumerate(tb.rows):
                key = row.cells[0].text.strip()
                if key == "Oferta de título doble":
                    # fila con SÍ / NO en la celda intermedia
                    mid = row.cells[1].text.strip().upper() if len(row.cells) > 2 else ""
                    val = str(ident.get("Oferta de título doble", "")).strip().upper()
                    if mid in ("SÍ", "SI", "NO") and mid.replace("Í", "I") == val.replace("Í", "I"):
                        set_cell_text(row.cells[-1], "X")
                    continue
                if key in ident:
                    set_cell_text(row.cells[-1], ident[key])
            break


def _year_cells(row):
    """Celdas-año distintas de una fila (las celdas combinadas se repiten en row.cells)."""
    out = []
    for c in row.cells:
        if YEAR_RE.match(c.text) and not any(o._tc is c._tc for o in out):
            out.append(c)
    return out


def relabel_year_headers(doc, data):
    """Sustituye los años de las cabeceras de la plantilla por `anios`/`anios_egresados`.

    Las tablas de egresados se reconocen por la cabecera "AÑO DE EGRESO".
    """
    general = data.get("anios") or []
    egresados = data.get("anios_egresados") or []
    for tb in doc.tables:
        is_egr = any("AÑO DE EGRESO" in c.text.upper() for row in tb.rows[:2] for c in row.cells)
        nuevos = egresados if is_egr else general
        if not nuevos:
            continue
        for row in tb.rows:
            cells = _year_cells(row)
            if not cells:
                continue
            if len(cells) != len(nuevos):
                print(f"AVISO: una tabla tiene {len(cells)} columnas de años y el JSON aporta "
                      f"{len(nuevos)}; no se han cambiado sus cabeceras.", file=sys.stderr)
                continue
            for c, year in zip(cells, nuevos):
                if c.text.strip() != year:
                    set_cell_text(c, year)


def check_scale(label, scope, year, value):
    """Avisa si un indicador de satisfacción (escala 0-5) supera el máximo."""
    if "satisfacci" not in label.lower():
        return
    try:
        num = float(str(value).replace(",", ".").split()[0])
    except (ValueError, IndexError):
        return
    if num > 5:
        print(f"AVISO: {label[:40]}… ({scope}, {year}) = {value} supera la escala 0-5; "
              "revisa el dato.", file=sys.stderr)


def check_extraction(data):
    """Avisa si falta la procedencia de los datos o si vienen del entorno de pruebas."""
    ext = data.get("extraccion") or {}
    entorno = str(ext.get("entorno", "")).lower()
    if not entorno or not ext.get("fecha"):
        print("AVISO: falta 'extraccion' (entorno y fecha) en el JSON; documenta de dónde y "
              "cuándo se extrajeron los indicadores.", file=sys.stderr)
    if "datapru" in entorno:
        print("AVISO: los indicadores proceden del entorno de PRUEBAS (datapru.uca.es); "
              "repite la extracción en producción (data.uca.es).", file=sys.stderr)


def fill_indicators(doc, data):
    relabel_year_headers(doc, data)
    check_extraction(data)
    inds = data.get("indicadores", {})
    if not inds:
        return 0
    written = 0
    for tb in doc.tables:
        scope = None
        colyears = {}  # col_index -> year
        first_year_col = None
        for row in tb.rows:
            cells = row.cells
            texts = [c.text.strip() for c in cells]
            rowtext = " ".join(texts)
            # ¿fila de cabecera con años?
            year_cols = {i: t for i, t in enumerate(texts) if YEAR_RE.match(t)}
            if year_cols:
                colyears = year_cols
                first_year_col = min(year_cols)
            # ¿indica ámbito? Solo en cabeceras: las etiquetas de datos dicen "Título-Centro".
            up = rowtext.upper()
            if not CODE_RE.search(up):
                if "CENTRO" in up and "TÍTULO" not in up and "TITULO" not in up:
                    scope = "centro"
                elif "TÍTULO" in up or "TITULO" in up:
                    scope = "titulo"
            # ¿fila de datos? El código ISGC debe ser el PREFIJO de la etiqueta,
            # para no confundir con filas "NOTA: ISGC-P07-05 al ISGC-P07-10 ...".
            if not (colyears and first_year_col is not None and scope):
                continue
            leading = distinct_leading(cells, first_year_col)
            label = leading[0] if leading else ""
            if label.upper().startswith("NOTA"):
                continue
            m = CODE_RE.match(label.strip())
            if not m:
                continue
            code = m.group(0)
            key = code
            if code == "ISGC-P04-09" and len(leading) >= 2:
                key = f"{code}|{leading[-1]}"
            values = inds.get(key)
            if not values:
                continue
            byyear = values.get(scope, {})
            for ci, year in colyears.items():
                if year in byyear and byyear[year] not in (None, ""):
                    check_scale(label, scope, year, byyear[year])
                    set_cell_text(tb.cell(row._index if hasattr(row, "_index") else _row_index(tb, row), ci), byyear[year])
                    written += 1
    return written


def _row_index(tb, row):
    for i, r in enumerate(tb.rows):
        if r._tr is row._tr:
            return i
    return 0


def fill_puntos(doc, data):
    puntos = data.get("puntos", {})
    if not puntos:
        return
    tables = doc.tables
    for crit, (fi, di) in PUNTOS_TABLES.items():
        info = puntos.get(str(crit)) or puntos.get(crit)
        if not info:
            continue
        # Fuertes
        fuertes = info.get("fuertes") or []
        if fuertes and fi < len(tables):
            tbf = tables[fi]
            if "fuerte" in tbf.cell(0, 0).text.lower() and len(tbf.rows) > 1:
                set_cell_text(tbf.cell(1, 0), "\n".join(f"- {x}" for x in fuertes))
        # Débiles
        debiles = info.get("debiles") or []
        if debiles and di < len(tables):
            tbd = tables[di]
            if "débil" in tbd.cell(0, 0).text.lower() or "debil" in tbd.cell(0, 0).text.lower():
                data_rows = [r for r in range(1, len(tbd.rows))]
                for j, deb in enumerate(debiles):
                    punto = deb.get("punto", "")
                    acciones = deb.get("acciones", [])
                    accion_txt = "\n".join(f"Acción de mejora nº {k+1}: {a}" for k, a in enumerate(acciones))
                    if j < len(data_rows):
                        r = data_rows[j]
                    else:
                        tbd.add_row()
                        r = len(tbd.rows) - 1
                    set_cell_text(tbd.cell(r, 0), f"Punto débil nº {j+1}: {punto}")
                    if len(tbd.columns) > 1:
                        set_cell_text(tbd.cell(r, 1), accion_txt)
                # limpia filas placeholder no usadas
                for r in data_rows[len(debiles):]:
                    for c in range(len(tbd.columns)):
                        set_cell_text(tbd.cell(r, c), "")


def fill_accua(doc, data):
    """Rellena la tabla de recomendaciones de ACCUA (Criterio 2).

    data['accua'] = {"fecha": "...", "items": [{"rec": "...", "accion": "..."}, ...]}
    Localiza la tabla por su cabecera ("Recomendaciones recibidas").
    """
    accua = data.get("accua")
    if not accua or not accua.get("items"):
        return
    fecha = accua.get("fecha", "")
    items = accua["items"]
    for tb in doc.tables:
        head = " ".join(c.text.lower() for c in tb.rows[0].cells)
        if "recomendaciones recibidas" in head:
            data_rows = list(range(1, len(tb.rows)))
            for j, it in enumerate(items):
                if j < len(data_rows):
                    r = data_rows[j]
                else:
                    tb.add_row()
                    r = len(tb.rows) - 1
                if len(tb.columns) >= 3:
                    set_cell_text(tb.cell(r, 0), fecha if j == 0 else "")
                    set_cell_text(tb.cell(r, 1), it.get("rec", ""))
                    set_cell_text(tb.cell(r, 2), it.get("accion", ""))
            for r in data_rows[len(items):]:
                for c in range(len(tb.columns)):
                    set_cell_text(tb.cell(r, c), "")
            break


def fill_plan_mejora(doc, data):
    plan = data.get("plan_mejora", [])
    if not plan:
        return
    tables = doc.tables
    if PLAN_MEJORA_TABLE >= len(tables):
        return
    tb = tables[PLAN_MEJORA_TABLE]
    if "prioridad" not in " ".join(c.text.lower() for c in tb.rows[0].cells):
        return
    data_rows = list(range(1, len(tb.rows)))
    for j, fila in enumerate(plan):
        if j < len(data_rows):
            r = data_rows[j]
        else:
            tb.add_row()
            r = len(tb.rows) - 1
        cols = tb.columns
        vals = [
            fila.get("origen", ""),
            f"Acción de mejora: {fila.get('accion','')}\nEvidencia: {fila.get('evidencia','')}",
            fila.get("prioridad", ""),
            fila.get("responsable", ""),
            fila.get("inicio", ""),
            fila.get("fin", ""),
            fila.get("indicador", ""),
        ]
        for ci, v in enumerate(vals):
            if ci < len(cols):
                set_cell_text(tb.cell(r, ci), v)
    # limpia filas placeholder no usadas
    for r in data_rows[len(plan):]:
        for ci in range(len(tb.columns)):
            set_cell_text(tb.cell(r, ci), "")


def main():
    ap = argparse.ArgumentParser(description="Rellena la plantilla del Autoinforme de Seguimiento (UCA).")
    ap.add_argument("--plantilla", required=True)
    ap.add_argument("--datos", required=True)
    ap.add_argument("--salida", required=True)
    args = ap.parse_args()

    with open(args.datos, encoding="utf-8") as f:
        data = json.load(f)

    doc = docx.Document(args.plantilla)
    fill_header(doc, data.get("encabezado"))
    fill_identity(doc, data.get("identificacion"))
    n = fill_indicators(doc, data)
    fill_accua(doc, data)
    fill_puntos(doc, data)
    fill_plan_mejora(doc, data)
    doc.save(args.salida)
    print(f"OK: {args.salida} | valores de indicadores escritos: {n}")


if __name__ == "__main__":
    main()
