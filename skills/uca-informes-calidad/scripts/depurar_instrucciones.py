#!/usr/bin/env python3
"""Depura el texto de instrucciones de la plantilla del autoinforme.

Elimina las instrucciones de cumplimentación y las directrices, para dejar un único
documento limpio tras redactar el análisis. Recorre el cuerpo y las celdas de las tablas
(también las anidadas). Criterios de borrado:
  - Párrafos cuyo texto va entero en ROJO (instrucciones) o AZUL (directrices ACCUA),
    según el convenio de colores de la plantilla.
  - Párrafos entre corchetes "[ ... ]" (instrucciones) o con marcadores "XXXX"/"xxxx".
  - En párrafos mixtos, solo los fragmentos (runs) en rojo o azul.
En las celdas, las etiquetas cortas solo en azul (encabezados de ejemplo como
«Suspensos» o «%») se conservan y pasan a negro: son parte de la estructura de la tabla.
Conserva el texto en negro (análisis), las tablas y los marcadores "[PENDIENTE: …]".
En una celda que se quedaría sin párrafos se vacía el texto en lugar de borrar el último.
Tras quitar los fragmentos de color de un párrafo mixto, borra el párrafo si solo queda
puntuación, recorta el guion final que queda colgando y lista el resto para revisarlo: el
texto en negro de un párrafo mixto puede ser parte de una instrucción o una opción que
no procede para el título, y eso solo puede decidirlo quien redacta.

Uso:
  python depurar_instrucciones.py --entrada relleno.docx --salida final.docx
"""
import argparse
import re
import sys

try:
    import docx
except ImportError:
    sys.exit("Falta python-docx: pip install python-docx --break-system-packages")

XXXX = re.compile(r"[Xx]{3,}")
PENDIENTE = "[PENDIENTE"
MAX_PALABRAS_ETIQUETA = 6
SOLO_PUNTUACION = re.compile(r"^[\s\-–—.:;,()]*$")
GUION_COLGANTE = re.compile(r"(?<=\s)[-–—]\s*\.?\s*$")
MAX_AVISOS = 40


def _rgb(run):
    try:
        c = run.font.color
        if c is not None and c.rgb is not None:
            return (c.rgb[0], c.rgb[1], c.rgb[2])
    except Exception:
        pass
    return None


def _is_reddish(rgb):
    r, g, b = rgb
    return r >= 120 and g < 90 and b < 90


def _is_bluish(rgb):
    r, g, b = rgb
    # Cubre el azul oscuro (0070C0) y el azul claro (00B0F0) de las directrices ACCUA.
    return b >= 120 and r < 90 and g < b


def _is_colored(run):
    rgb = _rgb(run)
    return bool(rgb) and (_is_reddish(rgb) or _is_bluish(rgb))


def is_instruction(p):
    t = p.text.strip()
    if not t or PENDIENTE in t:
        return False
    if t.startswith("[") or XXXX.search(t):
        return True
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(_is_colored(r) for r in runs)


def is_etiqueta_tabla(p):
    t = p.text.strip()
    if not t or t.startswith("[") or XXXX.search(t) or len(t.split()) > MAX_PALABRAS_ETIQUETA:
        return False
    colores = [_rgb(r) for r in p.runs if r.text.strip() and _is_colored(r)]
    return bool(colores) and all(_is_bluish(c) for c in colores)


def _borrar(p, en_celda):
    if en_celda and len(p._element.getparent().findall(p._element.tag)) == 1:
        for r in p.runs:
            r.text = ""
    else:
        p._element.getparent().remove(p._element)


def _recortar_final(p, patron):
    texto = "".join(r.text for r in p.runs)
    m = patron.search(texto)
    if not m:
        return
    sobran = len(texto) - m.start()
    for r in reversed(p.runs):
        if sobran <= 0:
            break
        n = min(len(r.text), sobran)
        r.text = r.text[:len(r.text) - n]
        sobran -= n


def depurar_parrafos(parrafos, en_celda=False, revisar=None):
    """Devuelve (párrafos borrados, fragmentos borrados).

    Añade a `revisar` el texto en negro que queda en los párrafos mixtos depurados."""
    borrados = fragmentos = 0
    for p in list(parrafos):
        if en_celda and is_etiqueta_tabla(p):
            for r in p.runs:
                r.font.color.rgb = None
        elif is_instruction(p):
            _borrar(p, en_celda)
            borrados += 1
        elif PENDIENTE not in p.text:
            color = [r for r in p.runs if r.text.strip() and _is_colored(r)]
            if not color:
                continue
            for r in color:
                r.text = ""
            fragmentos += len(color)
            if SOLO_PUNTUACION.match("".join(r.text for r in p.runs)):
                _borrar(p, en_celda)
                borrados += 1
                continue
            _recortar_final(p, GUION_COLGANTE)
            if revisar is not None:
                revisar.append(p.text.strip())
    return borrados, fragmentos


def depurar_tablas(tablas, revisar=None):
    borrados = fragmentos = 0
    for tabla in tablas:
        celdas = set()
        for fila in tabla.rows:
            for celda in fila.cells:
                if celda._tc in celdas:
                    continue
                celdas.add(celda._tc)
                b, f = depurar_parrafos(celda.paragraphs, en_celda=True, revisar=revisar)
                bt, ft = depurar_tablas(celda.tables, revisar)
                borrados, fragmentos = borrados + b + bt, fragmentos + f + ft
    return borrados, fragmentos


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--entrada", required=True)
    ap.add_argument("--salida", required=True)
    args = ap.parse_args()
    d = docx.Document(args.entrada)
    revisar = []
    borrados, fragmentos = depurar_parrafos(d.paragraphs, revisar=revisar)
    bt, ft = depurar_tablas(d.tables, revisar)
    d.save(args.salida)
    print(f"OK: {args.salida} | párrafos de instrucciones eliminados: {borrados + bt} "
          f"(en tablas: {bt}) | fragmentos en rojo/azul eliminados: {fragmentos + ft}")
    avisar_revision(revisar)


def avisar_revision(revisar):
    """Lista, sin repetir, los párrafos mixtos cuyo resto en negro hay que revisar."""
    unicos = list(dict.fromkeys(t for t in revisar if t))
    if not unicos:
        return
    print(f"REVISAR: {len(unicos)} párrafos mezclaban instrucciones y texto en negro. "
          "Comprueba que lo que queda no sea parte de una instrucción ni una opción que no "
          "procede para el título:", file=sys.stderr)
    for t in unicos[:MAX_AVISOS]:
        print(f"  - {t[:110]}{'…' if len(t) > 110 else ''}", file=sys.stderr)
    if len(unicos) > MAX_AVISOS:
        print(f"  … y {len(unicos) - MAX_AVISOS} más.", file=sys.stderr)


if __name__ == "__main__":
    main()
