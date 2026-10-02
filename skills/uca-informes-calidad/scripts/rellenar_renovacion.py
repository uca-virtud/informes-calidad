#!/usr/bin/env python3
"""Rellena la parte determinista del Autoinforme de Renovación de la Acreditación (UCA/ACCUA).

La renovación se presenta en la plataforma web de ACCUA; la Unidad de Calidad facilita
una plantilla .docx por criterio (assets/plantillas_renovacion/*CRITERIO_n*.docx) para
preparar el texto antes de volcarlo. Este script genera, sin alterar su formato:

  - Criterio_1.docx … Criterio_7.docx:
      * sin la línea interna "Formato para la redacción del Criterio n. GyM"
      * ficha de identificación (criterio 1) y denominación del título (criterios 2-7)
      * puntos fuertes agrupados por curso del periodo de acreditación
      * puntos débiles: curso | punto débil | propuestas de mejora | impacto
      * criterio 2: recomendaciones de la renovación/modificación anterior y de los
        informes de seguimiento de ACCUA del periodo
      * tablas "Nombre evidencia | Enlace evidencia" (criterio 5)
  - Anexo_I_Indicadores.docx: las tablas ISGC P01-P07 (TÍTULO y CENTRO) del Anexo 1 de
    la plantilla de seguimiento, rellenas con los indicadores.

Uso:
  python rellenar_renovacion.py --plantillas assets/plantillas_renovacion \
         --plantilla-indicadores assets/plantilla_autoinforme_seguimiento.docx \
         --datos datos.json --salida-dir salida/

El esquema de datos.json está en references/estructura-renovacion.md.
El análisis narrativo NO lo rellena este script: se redacta después sobre cada .docx.
"""
import argparse
import copy
import json
import os
import re
import sys

try:
    import docx
except ImportError:
    sys.exit("Falta python-docx. Instala con: pip install python-docx --break-system-packages")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rellenar_autoinforme import YEAR_RE, fill_indicators, set_cell_text  # noqa: E402

CRITERIO_RE = re.compile(r"CRITERIO_(\d)", re.IGNORECASE)
ANEXO_INICIO_RE = re.compile(r"^ANEXO\s+1\s*:", re.IGNORECASE)
ANEXO_FIN_RE = re.compile(r"^ANEXO\s+2\s*:", re.IGNORECASE)
FORMATO_RE = re.compile(r"^\s*Formato para la redacci[oó]n del Criterio", re.IGNORECASE)


def _header_text(tb, row=0):
    return " ".join(c.text.strip().lower() for c in tb.rows[row].cells)


def _fill_rows(tb, first_row, values):
    """Vuelca `values` (lista de listas) desde `first_row`.

    Clona la primera fila de datos cuando faltan filas, para conservar su formato, y
    elimina las filas de ejemplo sobrantes. Sin valores, deja la plantilla intacta.
    """
    if not values:
        return
    modelo = copy.deepcopy(tb.rows[first_row]._tr)
    while len(tb.rows) - first_row < len(values):
        tb.rows[-1]._tr.addnext(copy.deepcopy(modelo))
    for i, fila in enumerate(values):
        for c, valor in enumerate(fila):
            if c < len(tb.columns):
                set_cell_text(tb.cell(first_row + i, c), valor)
    for row in list(tb.rows)[first_row + len(values):]:
        row._tr.getparent().remove(row._tr)


def _texto(valor):
    if isinstance(valor, list):
        return "\n".join(str(v) for v in valor)
    return "" if valor is None else str(valor)


def _accion_evidencia(item):
    return f"Acción: {_texto(item.get('accion'))}\nEvidencia: {_texto(item.get('evidencia'))}"


def fill_ficha(doc, ident):
    """Rellena las filas 'etiqueta | valor' cuya etiqueta coincide con una clave de `ident`."""
    if not ident:
        return
    for tb in doc.tables:
        for row in tb.rows:
            cells = row.cells
            if len(cells) < 2 or cells[0]._tc is cells[-1]._tc:
                continue
            clave = cells[0].text.strip()
            if clave in ident and ident[clave] not in (None, ""):
                set_cell_text(cells[-1], ident[clave])


def fill_fuertes(doc, fuertes):
    """`fuertes`: {"2020/21": ["PF-…"], …} o lista simple de puntos fuertes."""
    if not fuertes:
        return
    if isinstance(fuertes, dict):
        lineas = []
        for curso, items in fuertes.items():
            lineas.append(f"{curso}:")
            lineas.extend(f"- {x}" for x in items)
    else:
        lineas = [f"- {x}" for x in fuertes]
    for tb in doc.tables:
        if len(tb.rows) >= 2 and "puntos fuertes" in tb.cell(0, 0).text.lower():
            set_cell_text(tb.cell(1, 0), "\n".join(lineas))
            return


def fill_debiles(doc, debiles):
    if not debiles:
        return
    filas = [[_texto(d.get("curso")), _texto(d.get("punto")),
              _texto(d.get("propuestas")), _texto(d.get("impacto"))] for d in debiles]
    for tb in doc.tables:
        if len(tb.rows) >= 3 and "puntos débiles" in tb.cell(0, 0).text.lower():
            _fill_rows(tb, 2, filas)
            return


def fill_recomendaciones(doc, data):
    """Tablas del subcriterio 2.4 (solo existen en la plantilla del criterio 2)."""
    renov = data.get("recomendaciones_renovacion") or []
    segui = data.get("recomendaciones_seguimiento") or []
    for tb in doc.tables:
        cab = _header_text(tb)
        if "informes de seguimiento" in cab:
            _fill_rows(tb, 1, [[_texto(it.get("convocatoria")), _texto(it.get("rec")),
                                _accion_evidencia(it), _texto(it.get("impacto"))] for it in segui])
        elif cab.startswith("recomendaciones") and len(tb.columns) == 3:
            _fill_rows(tb, 1, [[_texto(it.get("rec")), _accion_evidencia(it),
                                _texto(it.get("impacto"))] for it in renov])


def fill_evidencias(doc, evidencias):
    if not evidencias:
        return
    for tb in doc.tables:
        if "nombre evidencia" in _header_text(tb):
            _fill_rows(tb, 1, [[_texto(e.get("nombre")), _texto(e.get("enlace"))] for e in evidencias])
            return


def quitar_linea_formato(doc):
    """Elimina la línea interna 'Formato para la redacción del Criterio n. GyM'."""
    for p in list(doc.paragraphs):
        if FORMATO_RE.match(p.text):
            p._element.getparent().remove(p._element)


def rellenar_criterio(ruta, num, data):
    doc = docx.Document(ruta)
    quitar_linea_formato(doc)
    fill_ficha(doc, data.get("identificacion"))
    info = (data.get("puntos") or {}).get(str(num)) or {}
    fill_fuertes(doc, info.get("fuertes"))
    fill_debiles(doc, info.get("debiles"))
    if num == 2:
        fill_recomendaciones(doc, data)
    fill_evidencias(doc, (data.get("evidencias") or {}).get(str(num)))
    return doc


def _texto_parrafo(el):
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return "".join(t.text or "" for t in el.iter(f"{w}t")).strip()


def recortar_a_anexo_1(doc):
    """Deja en el documento solo el Anexo 1 (desde 'ANEXO 1:' hasta antes de 'ANEXO 2:')."""
    body = doc.element.body
    hijos = list(body)
    textos = [_texto_parrafo(el) if el.tag.endswith("}p") else "" for el in hijos]
    inicio = next((i for i, t in enumerate(textos) if ANEXO_INICIO_RE.match(t)), None)
    fin = next((i for i, t in enumerate(textos) if ANEXO_FIN_RE.match(t)), None)
    if inicio is None or fin is None or fin <= inicio:
        sys.exit("No se localiza el Anexo 1 en la plantilla de indicadores.")
    for el in hijos[:inicio] + hijos[fin:]:
        if not el.tag.endswith("}sectPr"):
            body.remove(el)
    for p in doc.paragraphs:
        if ANEXO_INICIO_RE.match(p.text.strip()):
            for r in p.runs:
                r.text = re.sub(r"ANEXO\s+1", "ANEXO I", r.text)
            break


def avisar_anios(doc, data):
    cabecera = {c.text.strip() for tb in doc.tables for row in tb.rows for c in row.cells
                if YEAR_RE.match(c.text)}
    esperados = set(data.get("anios") or []) | set(data.get("anios_egresados") or [])
    faltan = sorted(esperados - cabecera)
    if faltan:
        print(f"AVISO: la plantilla de indicadores no tiene columnas para {', '.join(faltan)}; "
              "esos valores no se han volcado.")


def generar_anexo(plantilla, data, salida):
    doc = docx.Document(plantilla)
    n = fill_indicators(doc, data)
    avisar_anios(doc, data)
    recortar_a_anexo_1(doc)
    doc.save(salida)
    return n


def main():
    ap = argparse.ArgumentParser(description="Rellena las plantillas del Autoinforme de Renovación de la Acreditación (UCA).")
    ap.add_argument("--plantillas", required=True, help="Carpeta con las plantillas *CRITERIO_n*.docx")
    ap.add_argument("--plantilla-indicadores", required=True,
                    help="Plantilla de seguimiento, de la que se toman las tablas del Anexo 1")
    ap.add_argument("--datos", required=True)
    ap.add_argument("--salida-dir", required=True)
    args = ap.parse_args()

    with open(args.datos, encoding="utf-8") as f:
        data = json.load(f)
    os.makedirs(args.salida_dir, exist_ok=True)

    generados = []
    for nombre in sorted(os.listdir(args.plantillas)):
        m = CRITERIO_RE.search(nombre)
        if not (m and nombre.lower().endswith(".docx")) or nombre.startswith("~$"):
            continue
        num = int(m.group(1))
        salida = os.path.join(args.salida_dir, f"Criterio_{num}.docx")
        rellenar_criterio(os.path.join(args.plantillas, nombre), num, data).save(salida)
        generados.append(salida)

    anexo = os.path.join(args.salida_dir, "Anexo_I_Indicadores.docx")
    n = generar_anexo(args.plantilla_indicadores, data, anexo)
    for ruta in generados + [anexo]:
        print(f"OK: {ruta}")
    print(f"Criterios generados: {len(generados)} | valores de indicadores escritos: {n}")


if __name__ == "__main__":
    main()
