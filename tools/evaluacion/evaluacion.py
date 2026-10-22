#!/usr/bin/env python3
"""Organiza y mide las ejecuciones de evaluación de la skill de autoinformes.

Uso (desde la raíz del repositorio):
  python tools/evaluacion/evaluacion.py enlaces <url> [--todos]
  python tools/evaluacion/evaluacion.py caso <fuentes.json> [--sin-descarga]
  python tools/evaluacion/evaluacion.py nueva <id> [--modelo-a M] [--modelo-b M]
  python tools/evaluacion/evaluacion.py medir <Enn>
  python tools/evaluacion/evaluacion.py valoracion <Enn>

Ver docs/PROCEDIMIENTO.md.
"""
import argparse
import csv
import json
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[2]
DOC_EVAL = RAIZ / "docs" / "evaluacion"
EVALS = RAIZ / "evals"
CASOS_DIR = EVALS / "casos"
EJEC_DIR = EVALS / "ejecuciones"
CASOS_CSV = EVALS / "casos.csv"
SKILL_DIR = RAIZ / "skills" / "uca-informes-calidad"
PLANTILLA_DOCX = SKILL_DIR / "assets" / "plantilla_autoinforme_seguimiento.docx"

MODELO_A = "claude-opus-5.5"
MODELO_B = "gpt-6.1-sol"
DIMENSIONES = ["D1", "D2", "D3", "D4", "D5", "D6", "D7"]
CAUSAS = ["S", "D", "I", "H", "E"]
ENTRADA_OBLIGATORIA = {"seguimiento": ["autoinforme-anterior", "memoria"]}
ROLES = {"autoinforme-humano": "referencia", "autoinforme-anterior": "entrada", "memoria": "entrada",
         "informe-accua": "entrada", "plan-mejora": "entrada"}
VALORES_CASO = {"nivel": {"grado", "máster"}, "tipo": {"seguimiento", "renovación"},
                "conjunto": {"aprendizaje", "reserva"}}
FORMATO_CASO = {"id": r"[A-Za-z0-9]+(-[A-Za-z0-9]+)*", "curso": r"\d{4}-\d{2}", "convocatoria": r"\d{4}"}
DOMINIOS_DESCARGA = ("uca.es", "juntadeandalucia.es")

RE_CODIGO = re.compile(r"\b(PF|PD|AM)[\s\-–]*[A-Z]{2,6}[\s\-–]*\d+(?:\.\d+)*\b")
RE_INDICADOR = re.compile(r"\bP0\d[\-_ ]?\d{2}")
RE_NUMERO = re.compile(r"(?<![\w.,])-?\d{1,3}(?:\.\d{3})*(?:,\d+)?(?:\.\d+)?(?![\w])")
MARCA_INI, MARCA_FIN = "<!-- medir:inicio -->", "<!-- medir:fin -->"
MARCA_VAL_INI, MARCA_VAL_FIN = "<!-- valoracion:inicio -->", "<!-- valoracion:fin -->"
SECCION_VALORACION = "## Valoración de la coordinación (completa `evaluacion.py valoracion`)"
# Aspecto de la hoja de valoración → claves que puntúa ("Datos" es de la hoja anterior a D1/D2 separados).
ASPECTOS_HOJA = {"identificación": ["D1"], "indicadores": ["D2"], "datos": ["D1", "D2"], "cobertura": ["D3"],
                 "análisis": ["D4"], "mejora": ["D5"], "recomendaciones": ["D6"], "redacción": ["D7"],
                 "esfuerzo": ["Esfuerzo"], "global": ["Global"]}
TEXTOS_HOJA = {"¿qué cambiaría del borrador?": "Qué cambiaría", "¿qué echa en falta?": "Qué echa en falta",
               "otros comentarios": "Otros comentarios"}
UTILIDAD = ["Esfuerzo", "Global"]


def error(msg):
    sys.exit("ERROR: " + msg)


def leer_casos():
    with open(CASOS_CSV, encoding="utf-8") as f:
        return {fila["id"]: fila for fila in csv.DictReader(f)}


def version_rubrica():
    m = re.search(r"^Versión:\s*(\S+)", (DOC_EVAL / "rubrica.md").read_text(encoding="utf-8"), re.M)
    return m.group(1) if m else "?"


def version_skill():
    try:
        commit = subprocess.run(["git", "log", "-1", "--format=%h", "--", str(SKILL_DIR)],
                                cwd=RAIZ, capture_output=True, text=True).stdout.strip()
        cambios = subprocess.run(["git", "status", "--porcelain", "--", str(SKILL_DIR)],
                                 cwd=RAIZ, capture_output=True, text=True).stdout.strip()
    except OSError:
        return "desconocida"
    return (commit or "sin commit") + (" + cambios sin commit" if cambios else "")


# --- nueva -------------------------------------------------------------------

def comprobar_caso(caso):
    base = CASOS_DIR / caso["id"]
    entrada, referencia = base / "entrada", base / "referencia"
    faltan = [n + ".*" for n in ENTRADA_OBLIGATORIA.get(caso["tipo"], []) if not list(entrada.glob(n + ".*"))]
    if not entrada.is_dir() or not any(entrada.iterdir()):
        faltan.append("entrada/ (vacía)")
    if not list(referencia.glob("autoinforme-humano.*")):
        faltan.append("referencia/autoinforme-humano.*")
    if faltan:
        error(f"al caso {caso['id']} le faltan: {', '.join(faltan)} (en {base})")


def siguiente_numero():
    nums = [int(m.group(1)) for m in (re.match(r"E(\d+)-", p.name) for p in EJEC_DIR.glob("E*-*")) if m]
    return max(nums, default=0) + 1


def cmd_nueva(args):
    casos = leer_casos()
    caso = casos.get(args.id) or error(f"el caso {args.id} no está en evals/casos.csv")
    if caso["conjunto"] == "excluido":
        error(f"{args.id} está excluido: la skill usa textos suyos como ejemplos")
    comprobar_caso(caso)
    nombre = f"E{siguiente_numero():02d}-{caso['id']}"
    carpeta = EJEC_DIR / nombre
    carpeta.mkdir(parents=True)
    accua = caso.get("seguimiento_accua") or "no consta: se asume sí"
    valores = {
        "<Enn-id>": nombre, "<Enn>": nombre.split("-")[0], "<id>": caso["id"],
        "<titulo>": caso["titulo"], "<tipo>": caso["tipo"], "<curso>": caso["curso"],
        "<convocatoria>": caso["convocatoria"], "<sí / no / no consta: se asume sí>": accua,
        "<aprendizaje / reserva>": caso["conjunto"], "<AAAA-MM-DD>": date.today().isoformat(),
        "<commit de skills/ y si hay cambios sin commit>": version_skill(),
        "<n>": version_rubrica(), "<modelo>": None,
    }
    texto = (DOC_EVAL / "plantilla-ficha.md").read_text(encoding="utf-8")
    for clave, valor in valores.items():
        if valor is not None:
            texto = texto.replace(clave, valor, 1 if clave == "<AAAA-MM-DD>" else -1)
    texto = texto.replace("<modelo>", args.modelo_a, 1).replace("<modelo>", args.modelo_b, 1)
    (carpeta / "ficha.md").write_text(texto, encoding="utf-8")
    print(f"Creada {carpeta.relative_to(RAIZ)}")
    print(f"  entrada:    {(CASOS_DIR / caso['id'] / 'entrada').relative_to(RAIZ)}")
    print(f"  referencia: {(CASOS_DIR / caso['id'] / 'referencia').relative_to(RAIZ)} (no se da a A al generar)")


# --- caso ----------------------------------------------------------------------

def leer_manifiesto(ruta, sin_descarga):
    try:
        man = json.loads(Path(ruta).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        error(f"no se puede leer el manifiesto {ruta}: {e}")
    with open(CASOS_CSV, encoding="utf-8") as f:
        columnas = next(csv.reader(f))
    faltan = [c for c in columnas if c != "seguimiento_accua" and not str(man.get(c, "")).strip()]
    if faltan:
        error(f"faltan campos en el manifiesto: {', '.join(faltan)}")
    for campo, permitidos in VALORES_CASO.items():
        if man[campo] not in permitidos:
            error(f"{campo} = «{man[campo]}»; valores admitidos: {', '.join(sorted(permitidos))}")
    for campo, patron in FORMATO_CASO.items():
        if not re.fullmatch(patron, man[campo]):
            error(f"{campo} = «{man[campo]}» no tiene el formato esperado ({patron})")
    if man["id"] in leer_casos():
        error(f"el caso {man['id']} ya está en evals/casos.csv")
    documentos = man.get("documentos", [])
    roles = [d.get("rol") for d in documentos]
    desconocidos = [r for r in roles if r not in ROLES]
    if desconocidos:
        error(f"roles desconocidos: {desconocidos}; admitidos: {', '.join(ROLES)}")
    if len(set(roles)) != len(roles):
        error("hay roles repetidos en documentos: elige un único documento por rol")
    if not sin_descarga:
        necesarios = ["autoinforme-humano"] + ENTRADA_OBLIGATORIA.get(man["tipo"], [])
        ausentes = [r for r in necesarios if r not in roles]
        if ausentes:
            error(f"faltan documentos obligatorios en el manifiesto: {', '.join(ausentes)}")
        for d in documentos:
            comprobar_url(d.get("url", ""))
    return man, columnas, documentos


def comprobar_url(url):
    from urllib.parse import urlparse
    partes = urlparse(url)
    host = partes.hostname or ""
    if partes.scheme != "https" or not any(host == d or host.endswith("." + d) for d in DOMINIOS_DESCARGA):
        error(f"URL no admitida (solo https de {', '.join(DOMINIOS_DESCARGA)}): {url}")


def leer_url(url):
    from urllib.parse import quote
    from urllib.request import Request, urlopen
    comprobar_url(url)
    peticion = Request(quote(url, safe=":/?&=%#~+@,;"), headers={"User-Agent": "uca-informes-calidad/evaluacion"})
    try:
        with urlopen(peticion, timeout=60) as r:
            comprobar_url(r.geturl())
            return r.read(), r.headers.get_content_type(), r.headers.get_content_charset() or "utf-8", r.geturl()
    except OSError as e:
        error(f"no se pudo descargar {url}: {e}")


def descargar(url):
    datos, tipo, _, _ = leer_url(url)
    if datos.startswith(b"%PDF"):
        return datos, ".pdf"
    if datos.startswith(b"PK") and ("wordprocessingml" in tipo or url.lower().endswith(".docx")):
        return datos, ".docx"
    error(f"{url} no es un PDF ni un DOCX (tipo {tipo}); busca el enlace directo al fichero")


def avisos_pdf(ruta):
    try:
        import pypdf
        lector = pypdf.PdfReader(str(ruta))
        muestra = "".join((p.extract_text() or "") for p in lector.pages[:3])
    except Exception as e:  # pypdf lanza excepciones propias muy variadas
        return [f"{ruta.name}: no se pudo leer con pypdf ({e})"]
    if len(muestra.strip()) < 200:
        return [f"{ruta.name}: apenas tiene texto extraíble (¿escaneado?)"]
    return []


def escribir_fuentes(base, man, filas):
    lineas = [f"# Fuentes del caso {man['id']}", "",
              f"{man['titulo']} · {man['centro']} · {man['tipo']} {man['curso']} (convocatoria {man['convocatoria']})",
              "", "| Rol | Fichero | URL | Página de origen | Descarga | Tamaño | SHA-256 |",
              "|---|---|---|---|---|---|---|"]
    lineas += ["| " + " | ".join(f) + " |" for f in filas]
    if man.get("notas"):
        lineas += ["", "## Notas", "", man["notas"]]
    (base / "fuentes.md").write_text("\n".join(lineas) + "\n", encoding="utf-8")


def anadir_fila_csv(man, columnas):
    ruta = CASOS_CSV
    contenido = ruta.read_text(encoding="utf-8")
    with open(ruta, "a", encoding="utf-8", newline="") as f:
        if contenido and not contenido.endswith("\n"):
            f.write("\n")
        csv.DictWriter(f, fieldnames=columnas, lineterminator="\n").writerow(
            {c: str(man.get(c, "")).strip() for c in columnas})


RE_ENLACE_DOC = re.compile(r"\.(pdf|docx?)(\?|$)|/wp-content/uploads/", re.I)
RE_ENLACE_CALIDAD = re.compile(r"autoinform|informe|memoria|plan[- ]de[- ]mejora|calidad|seguimiento|"
                               r"acreditaci|verificaci|modificaci|datos[- ]generales", re.I)
RE_ENLACE_SOCIAL = re.compile(r"^https?://([^/]+\.)?(facebook\.com|twitter\.com|x\.com|linkedin\.com|wa\.me)/", re.I)


def enlaces_de(html_texto, base_url):
    from html.parser import HTMLParser
    from urllib.parse import urljoin

    class Lector(HTMLParser):
        def __init__(self):
            super().__init__()
            self.enlaces, self.actual = [], None

        def handle_starttag(self, tag, attrs):
            href = dict(attrs).get("href")
            if tag == "a" and href and not href.startswith(("#", "mailto:", "javascript:")):
                self.actual = [urljoin(base_url, href), []]

        def handle_data(self, data):
            if self.actual:
                self.actual[1].append(data)

        def handle_endtag(self, tag):
            if tag == "a" and self.actual:
                self.enlaces.append((self.actual[0], " ".join(" ".join(self.actual[1]).split())))
                self.actual = None

    lector = Lector()
    lector.feed(html_texto)
    return list(dict.fromkeys(lector.enlaces))


def cmd_enlaces(args):
    datos, tipo, juego, url_final = leer_url(args.url)
    if tipo != "text/html":
        error(f"{args.url} no es una página HTML (tipo {tipo})")
    for url, texto in enlaces_de(datos.decode(juego, errors="replace"), url_final):
        if RE_ENLACE_SOCIAL.match(url):
            continue
        es_doc = bool(RE_ENLACE_DOC.search(url))
        if args.todos or es_doc or RE_ENLACE_CALIDAD.search(url + " " + texto):
            print(f"{'DOC' if es_doc else 'PÁG'} | {texto or '(sin texto)'} | {url}")


def cmd_caso(args):
    import hashlib
    man, columnas, documentos = leer_manifiesto(args.manifiesto, args.sin_descarga)
    base = CASOS_DIR / man["id"]
    if args.sin_descarga:
        comprobar_caso(man)
    hoy = date.today().isoformat()
    descargas = []
    if not args.sin_descarga:
        for d in documentos:
            carpeta = base / ROLES[d["rol"]]
            if list(carpeta.glob(d["rol"] + ".*")):
                error(f"ya existe {carpeta.relative_to(RAIZ)}/{d['rol']}.*; no se sobrescribe")
        for d in documentos:
            print(f"Descargando {d['rol']}: {d['url']}")
            descargas.append((d, *descargar(d["url"])))
    filas, avisos = [], []
    for d, datos, ext in descargas:
        ruta = base / ROLES[d["rol"]] / (d["rol"] + ext)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        ruta.write_bytes(datos)
        if ext == ".pdf":
            avisos += avisos_pdf(ruta)
        filas.append([d["rol"], str(ruta.relative_to(base)), d["url"], d.get("pagina", ""), hoy,
                      f"{len(datos) / 1024:.0f} KB", hashlib.sha256(datos).hexdigest()])
    if args.sin_descarga:
        for rol, sub in ROLES.items():
            for ruta in sorted((base / sub).glob(rol + ".*")):
                datos = ruta.read_bytes()
                filas.append([rol, str(ruta.relative_to(base)), "aportado a mano", "", hoy,
                              f"{len(datos) / 1024:.0f} KB", hashlib.sha256(datos).hexdigest()])
    base.mkdir(parents=True, exist_ok=True)
    escribir_fuentes(base, man, filas)
    comprobar_caso(man)
    anadir_fila_csv(man, columnas)
    print(f"Añadido {man['id']} a evals/casos.csv y creado {(base / 'fuentes.md').relative_to(RAIZ)}")
    for aviso in avisos:
        print("AVISO: " + aviso)
    print(f"Caso listo: python tools/evaluacion/evaluacion.py nueva {man['id']}")


# --- extracción de texto ------------------------------------------------------

def texto_docx(ruta):
    import docx
    doc = docx.Document(str(ruta))
    partes = [p.text for p in doc.paragraphs]
    vistas = set()
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                if celda._tc not in vistas:
                    vistas.add(celda._tc)
                    partes.append(celda.text)
    return "\n".join(partes)


def texto_pdf(ruta):
    import pypdf
    return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(str(ruta)).pages)


def texto(ruta):
    return texto_docx(ruta) if ruta.suffix.lower() == ".docx" else texto_pdf(ruta)


def a_numero(token):
    t = token.replace(".", "") if "," in token else token
    try:
        return float(t.replace(",", "."))
    except ValueError:
        return None


def es_anio(x):
    return x.is_integer() and 1990 <= x <= 2035


def valores_anexo(ruta_docx):
    """Valores numéricos de las filas de indicadores ISGC (P0x-yy) del borrador."""
    import docx
    valores = []
    for tabla in docx.Document(str(ruta_docx)).tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            if not celdas or not RE_INDICADOR.search(celdas[0]):
                continue
            vistas = set()
            for c in fila.cells[1:]:
                if c._tc in vistas:
                    continue
                vistas.add(c._tc)
                for tok in RE_NUMERO.findall(c.text):
                    x = a_numero(tok)
                    if x is not None and not es_anio(x):
                        valores.append(x)
    return valores


def metricas_documento(txt):
    return {
        "palabras": len(txt.split()),
        "pf": len({m.group(0) for m in RE_CODIGO.finditer(txt) if m.group(1) == "PF"}),
        "pd": len({m.group(0) for m in RE_CODIGO.finditer(txt) if m.group(1) == "PD"}),
        "am": len({m.group(0) for m in RE_CODIGO.finditer(txt) if m.group(1) == "AM"}),
        "pendientes": len(re.findall(r"\[PENDIENTE", txt, re.I)),
    }


def coincidencia_anexo(valores, txt_ref):
    ref = [a_numero(t) for t in RE_NUMERO.findall(txt_ref)]
    ref1 = {round(x, 1) for x in ref if x is not None}
    ref2 = {round(x, 2) for x in ref if x is not None}
    if not valores:
        return None
    hits = sum(1 for x in valores if round(x, 2) in ref2 or round(x, 1) in ref1)
    return {"valores": len(valores), "encontrados": hits, "porcentaje": round(100 * hits / len(valores), 1)}


# --- huellas de casos en la skill -------------------------------------------

def tejas(txt, n=8):
    palabras = re.findall(r"\w+", txt.lower())
    return {" ".join(palabras[i:i + n]) for i in range(len(palabras) - n + 1)}


def huellas_en_skill(casos):
    """Frases de 8 palabras que solo aparecen en la referencia de un título y también en la skill.

    Las frases presentes en la plantilla oficial o en las referencias de varios títulos se
    consideran institucionales y no se avisan."""
    genericas = tejas(texto_docx(PLANTILLA_DOCX)) if PLANTILLA_DOCX.exists() else set()
    por_caso = {}
    for caso in casos.values():
        for ref in (CASOS_DIR / caso["id"] / "referencia").glob("autoinforme-humano.*"):
            por_caso[caso["id"]] = tejas(texto(ref)) - genericas
    especificas = {}
    for caso_id, frases in por_caso.items():
        otras = set().union(*(f for c, f in por_caso.items() if c != caso_id))
        for frase in frases - otras:
            if sum(not w.isdigit() for w in frase.split()) >= 5:
                especificas[frase] = caso_id
    nombres = [c["titulo"] for c in casos.values()]
    hallazgos = []
    for f in sorted(SKILL_DIR.rglob("*")):
        if f.suffix not in {".md", ".json", ".py"} or "__pycache__" in f.parts:
            continue
        contenido = f.read_text(encoding="utf-8", errors="ignore")
        comunes = sorted(tejas(contenido) & especificas.keys())
        titulos = [n for n in nombres if n.lower() in contenido.lower()]
        if comunes or titulos:
            hallazgos.append({"fichero": str(f.relative_to(RAIZ)), "frases": len(comunes),
                              "casos": sorted({especificas[c] for c in comunes}),
                              "ejemplo": comunes[0] if comunes else "", "titulos": titulos})
    return hallazgos


# --- cotejos ------------------------------------------------------------------

def filas_tabla(md, seccion):
    m = re.search(rf"^## {seccion}\s*$(.*?)(?=^## |\Z)", md, re.M | re.S)
    if not m:
        return []
    filas = []
    for linea in m.group(1).splitlines():
        if linea.startswith("|") and not re.match(r"^\|\s*-", linea):
            filas.append([c.strip() for c in linea.strip().strip("|").split("|")])
    return filas[1:]


def version_cotejo(ruta):
    if ruta.exists():
        m = re.search(r"^- Versión de la rúbrica:\s*(\S+)", ruta.read_text(encoding="utf-8"), re.M)
        if m:
            return m.group(1)
    return "—"


def leer_cotejo(ruta):
    if not ruta.exists():
        return None
    md = ruta.read_text(encoding="utf-8")
    notas = {}
    for f in filas_tabla(md, "Puntuación"):
        if f and f[0] in DIMENSIONES and len(f) > 1:
            notas[f[0]] = int(f[1]) if f[1] in {"0", "1", "2"} else None
    bloqueantes = [f[0] for f in filas_tabla(md, "Bloqueantes") if len(f) > 1 and f[1].lower().startswith("s")]
    causas = {c: 0 for c in CAUSAS}
    pendientes = []
    for f in filas_tabla(md, "Discrepancias"):
        if len(f) > 4 and f[4].upper() in causas:
            causas[f[4].upper()] += 1
            if f[4].upper() in {"H", "I"}:
                pendientes.append(f"Discrepancia {f[0]} ({f[4].upper()}): {f[1]}")
    lecciones = [f[0] for f in filas_tabla(md, "Lecciones propuestas") if f and f[0] and not f[0].startswith("<")]
    return {"notas": notas, "bloqueantes": bloqueantes, "causas": causas,
            "pendientes": pendientes, "lecciones": lecciones}


def consolidar(a, b):
    final, pendientes = {}, []
    for d in DIMENSIONES:
        na = a["notas"].get(d) if a else None
        nb = b["notas"].get(d) if b else None
        final[d] = nb if nb is not None else na
        if na is not None and nb is not None and abs(na - nb) >= 2:
            pendientes.append(f"{d}: A={na}, B={nb}")
    validas = [v for v in final.values() if v is not None]
    media = round(sum(validas) / len(validas), 2) if validas else None
    return final, media, pendientes


# --- medir --------------------------------------------------------------------

def buscar_ejecucion(clave):
    candidatas = [p for p in EJEC_DIR.glob(f"{clave}*") if p.is_dir()]
    if len(candidatas) != 1:
        error(f"no hay una única ejecución que empiece por {clave}")
    return candidatas[0]


def nota(valor):
    return "—" if valor is None else valor


def tabla_md(cabecera, filas):
    lineas = ["| " + " | ".join(cabecera) + " |", "|" + "---|" * len(cabecera)]
    lineas += ["| " + " | ".join(str(c) for c in f) + " |" for f in filas]
    return "\n".join(lineas)


def bloque_resultado(r):
    m, h = r["borrador"], r["referencia"]
    ratio = f"{m['palabras']:,} / {h['palabras']:,} ({round(100 * m['palabras'] / h['palabras'])} %)".replace(",", ".")
    anexo = r["anexo"]
    partes = [
        tabla_md(["Métrica", "Borrador", "Humano"], [
            ["Palabras", m["palabras"], h["palabras"]], ["PF distintos", m["pf"], h["pf"]],
            ["PD distintos", m["pd"], h["pd"]], ["AM distintas", m["am"], h["am"]],
            ["`[PENDIENTE]`", m["pendientes"], h["pendientes"]]]),
        f"\nExtensión borrador / humano: {ratio}.",
        f"Valores del Anexo I presentes en algún lugar de la referencia (cota superior, no mide la celda): "
        + (f"{anexo['encontrados']}/{anexo['valores']} ({anexo['porcentaje']} %)." if anexo else "sin tablas de indicadores."),
        "",
        tabla_md(["Dim.", "A", "B", "Final"],
                 [[d, nota(r["notas_a"].get(d)), nota(r["notas_b"].get(d)), nota(r["final"][d])] for d in DIMENSIONES]),
        f"\nMedia final: {nota(r['media'])}. Bloqueantes: {', '.join(r['bloqueantes']) or 'ninguno'}. "
        f"Discrepancias por causa (B): " + ", ".join(f"{k}={v}" for k, v in r["causas_b"].items()) + ".",
    ]
    if r["huellas"]:
        partes.append("\n**Aviso de huellas de casos en la skill** (revisar que no sean datos de títulos):")
        partes += [f"- `{x['fichero']}`: {x['frases']} frases de {', '.join(x['casos']) or '—'}"
                   + (f", títulos {x['titulos']}" if x["titulos"] else "")
                   + (f" (p. ej. «{x['ejemplo']}»)" if x["ejemplo"] else "") for x in r["huellas"]]
    return "\n".join(str(p) for p in partes)


def actualizar_ficha(carpeta, r):
    ficha = carpeta / "ficha.md"
    if not ficha.exists():
        return
    txt = ficha.read_text(encoding="utf-8")
    txt = re.sub(re.escape(MARCA_INI) + ".*?" + re.escape(MARCA_FIN),
                 lambda _: f"{MARCA_INI}\n{bloque_resultado(r)}\n{MARCA_FIN}", txt, flags=re.S)
    pend = r["pendientes"] + pendientes_valoracion(carpeta) or ["Nada pendiente."]
    txt = re.sub(r"(## Pendiente de revisión experta\n\n).*?(?=\n## )",
                 lambda m: m.group(1) + "\n".join(f"- {p}" for p in pend) + "\n", txt, flags=re.S)
    ficha.write_text(txt, encoding="utf-8")


def campo_ficha(carpeta, campo):
    ficha = carpeta / "ficha.md"
    if ficha.exists():
        m = re.search(rf"^\| {re.escape(campo)} \| (.*?) \|", ficha.read_text(encoding="utf-8"), re.M)
        if m and not m.group(1).startswith("<"):
            return m.group(1)
    return "—"


def actualizar_historial():
    filas = []
    for carpeta in sorted(EJEC_DIR.glob("E*-*")):
        f = carpeta / "metricas.json"
        if not f.exists():
            continue
        r = json.loads(f.read_text(encoding="utf-8"))
        b, h = r["borrador"], r["referencia"]
        v = leer_json(carpeta / "valoracion.json")
        coord = v["media"] if v else None
        dif = round(coord - r["media"], 2) if v and coord is not None and r["media"] is not None else None
        filas.append([carpeta.name, campo_ficha(carpeta, "Fecha"), campo_ficha(carpeta, "Versión de la skill"),
                      campo_ficha(carpeta, "Modelo A (generador y autocotejo)"), r["rubrica"]] + [nota(r["final"][d]) for d in DIMENSIONES]
                     + [nota(r["media"]), ", ".join(r["bloqueantes"]) or "—",
                        f"{round(100 * b['palabras'] / h['palabras'])} %", f"{b['pd']} / {h['pd']}",
                        b["pendientes"], f"{r['anexo']['porcentaje']} %" if r["anexo"] else "—",
                        r["causas_b"]["S"], len(r["pendientes"]) + len(v["pendientes"] if v else []),
                        nota(coord), nota(dif)] + [nota(v["utilidad"].get(u) if v else None) for u in UTILIDAD])
    cab = ["Ejecución", "Fecha", "Skill", "Modelo A", "Rúbrica"] + DIMENSIONES + [
        "Media", "Bloq.", "Extensión", "PD IA/hum.", "[PEND.]", "Anexo I (cota sup.)", "Causas S", "Pend. revisión",
        "Coord.", "Δ coord.−IA"] + [f"{u} (1-5)" for u in UTILIDAD]
    bloque = (f"{MARCA_INI}\nGenerada por `evaluacion.py medir` y `valoracion`; no editar a mano. "
              f"«Coord.» es la media de la coordinación pasada a 0-2.\n\n{tabla_md(cab, filas)}\n{MARCA_FIN}")
    historial = EVALS / "historial.md"
    txt = historial.read_text(encoding="utf-8") if historial.exists() else ""
    if MARCA_INI in txt:
        txt = re.sub(re.escape(MARCA_INI) + ".*?" + re.escape(MARCA_FIN), lambda _: bloque, txt, flags=re.S)
    else:
        txt = f"# Historial de ejecuciones\n\n## Tabla de ejecuciones\n\n{bloque}\n"
    historial.write_text(txt, encoding="utf-8")


def cmd_medir(args):
    carpeta = buscar_ejecucion(args.ejecucion)
    caso_id = carpeta.name.split("-", 1)[1]
    casos = leer_casos()
    if caso_id not in casos:
        error(f"el caso {caso_id} no está en casos.csv")
    borrador = carpeta / "borrador.docx"
    refs = list((CASOS_DIR / caso_id / "referencia").glob("autoinforme-humano.*"))
    if not borrador.exists() or not refs:
        error("faltan borrador.docx o la referencia del caso")
    txt_b, txt_h = texto_docx(borrador), texto(refs[0])
    a, b = leer_cotejo(carpeta / "cotejo-A.md"), leer_cotejo(carpeta / "cotejo-B.md")
    if b is None:
        print("AVISO: no hay cotejo-B.md; la nota final será la de A.")
    final, media, pend_dim = consolidar(a, b)
    rubrica = version_cotejo(carpeta / ("cotejo-B.md" if b else "cotejo-A.md"))
    if rubrica not in ("—", version_rubrica()):
        print(f"AVISO: el cotejo usa la rúbrica {rubrica} y la vigente es la {version_rubrica()}; "
              "vuelve a puntuar con B.")
    bloqueantes = sorted(set((a or {}).get("bloqueantes", [])) | set((b or {}).get("bloqueantes", [])))
    pendientes = pend_dim + [f"Bloqueante {x}" for x in bloqueantes]
    for c in (a, b):
        for p in (c or {}).get("pendientes", []):
            if p not in pendientes:
                pendientes.append(p)
    r = {
        "rubrica": rubrica,
        "borrador": metricas_documento(txt_b), "referencia": metricas_documento(txt_h),
        "anexo": coincidencia_anexo(valores_anexo(borrador), txt_h),
        "notas_a": (a or {}).get("notas", {}), "notas_b": (b or {}).get("notas", {}),
        "final": final, "media": media, "bloqueantes": bloqueantes,
        "causas_b": (b or a or {}).get("causas", {c: 0 for c in CAUSAS}),
        "lecciones": sorted(set((a or {}).get("lecciones", [])) | set((b or {}).get("lecciones", []))),
        "pendientes": pendientes, "huellas": huellas_en_skill(casos),
    }
    (carpeta / "metricas.json").write_text(json.dumps(r, ensure_ascii=False, indent=2), encoding="utf-8")
    actualizar_ficha(carpeta, r)
    actualizar_historial()
    print(bloque_resultado(r))
    print(f"\nActualizados {carpeta.name}/metricas.json, ficha.md y evals/historial.md")


# --- valoración de la coordinación -----------------------------------------------

def leer_json(ruta):
    return json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else None


def pendientes_valoracion(carpeta):
    v = leer_json(carpeta / "valoracion.json")
    return v["pendientes"] if v else []


def a_escala_rubrica(nota_1_5):
    return round((nota_1_5 - 1) / 2, 2)


def decimal(x):
    if x is None:
        return "—"
    return str(int(x) if float(x).is_integer() else x).replace(".", ",")


def media(valores):
    validos = [x for x in valores if x is not None]
    return round(sum(validos) / len(validos), 2) if validos else None


def leer_hoja_valoracion(ruta):
    import openpyxl
    ws = openpyxl.load_workbook(ruta, data_only=True).active
    hoja = {"fichero": ruta.name, "rol": ws["C5"].value, "fecha": str(ws["C6"].value or "") or None,
            "notas": {}, "comentarios": {}, "graves": None, "comentario_graves": None, "textos": {}}
    formato_anterior = False
    for fila in range(1, ws.max_row + 1):
        etiqueta = str(ws.cell(fila, 2).value or "").strip().lower()
        texto_a = str(ws.cell(fila, 1).value or "").strip().lower()
        valor, comentario = ws.cell(fila, 4).value, ws.cell(fila, 5).value
        if etiqueta == "errores graves":
            respuesta = str(valor or "").strip().lower()
            hoja["graves"] = {"sí": True, "si": True, "no": False}.get(respuesta)
            hoja["comentario_graves"] = comentario
        elif etiqueta in ASPECTOS_HOJA:
            formato_anterior |= etiqueta == "datos"
            if valor not in (None, "") and not (isinstance(valor, (int, float)) and 1 <= valor <= 5 and valor == int(valor)):
                print(f"AVISO: {ruta.name}, fila {fila}: «{valor}» no es una nota de 1 a 5; se ignora.")
                valor = None
            for clave in ASPECTOS_HOJA[etiqueta]:
                hoja["notas"][clave] = None if valor in (None, "") else int(valor)
                if comentario:
                    hoja["comentarios"][clave] = str(comentario).strip()
        elif texto_a in TEXTOS_HOJA:
            contenido = ws.cell(fila + 1, 1).value
            if contenido:
                hoja["textos"][TEXTOS_HOJA[texto_a]] = str(contenido).strip()
    if not hoja["notas"]:
        error(f"{ruta.name} no tiene el formato de la hoja de valoración")
    if formato_anterior:
        print(f"AVISO: {ruta.name} usa la hoja anterior: «Datos» puntúa D1 y D2 a la vez; no hay D6 ni errores graves.")
    return hoja


def consolidar_valoracion(hojas, metricas):
    claves = DIMENSIONES + UTILIDAD
    notas_1_5 = {k: media([h["notas"].get(k) for h in hojas]) for k in claves}
    notas_0_2 = {d: None if notas_1_5[d] is None else a_escala_rubrica(notas_1_5[d]) for d in DIMENSIONES}
    respuestas = [h["graves"] for h in hojas if h["graves"] is not None]
    graves = None if not respuestas else any(respuestas)
    pendientes = []
    for d in DIMENSIONES:
        ia = metricas["final"].get(d)
        if notas_0_2[d] is not None and ia is not None and abs(notas_0_2[d] - ia) >= 1:
            pendientes.append(f"Coordinación {d}: coordinación={decimal(notas_0_2[d])}, IA={decimal(ia)}")
    graves_ia = [b for b in metricas["bloqueantes"] if b in ("B1", "B2")]
    if graves and not graves_ia:
        pendientes.append("Coordinación: señala errores graves y la IA no anotó B1 ni B2")
    elif graves is False and graves_ia:
        pendientes.append(f"Coordinación: no aprecia errores graves y la IA anotó {', '.join(graves_ia)}")
    return {"hojas": hojas, "notas_1_5": {d: notas_1_5[d] for d in DIMENSIONES}, "notas": notas_0_2,
            "media": media(notas_0_2.values()), "utilidad": {u: notas_1_5[u] for u in UTILIDAD},
            "graves": graves, "pendientes": pendientes}


def bloque_valoracion(v, metricas):
    fmt = decimal
    filas = []
    for d in DIMENSIONES:
        c, ia = v["notas"][d], metricas["final"].get(d)
        dif = None if c is None or ia is None else round(c - ia, 2)
        filas.append([d, fmt(v["notas_1_5"][d]), fmt(c), fmt(ia), fmt(dif)])
    dif_media = None if v["media"] is None or metricas["media"] is None else round(v["media"] - metricas["media"], 2)
    graves = {True: "sí", False: "no", None: "sin respuesta"}[v["graves"]]
    partes = [
        "Valoraciones: " + "; ".join(f"`{h['fichero']}` ({h['rol'] or 'rol sin indicar'}, {h['fecha'] or 'sin fecha'})"
                                     for h in v["hojas"]) + ".",
        "",
        tabla_md(["Dim.", "Coordinación (1-5)", "Equivale (0-2)", "IA (final)", "Diferencia"], filas),
        f"\nMedia: coordinación {fmt(v['media'])}, IA {fmt(metricas['media'])} (diferencia {fmt(dif_media)}). "
        + " ".join(f"{u} (1-5): {fmt(v['utilidad'][u])}." for u in UTILIDAD)
        + f" Errores graves según la coordinación: {graves}.",
    ]
    for h in v["hojas"]:
        detalle = [f"- {k}: {t}" for k, t in h["comentarios"].items()]
        if h["comentario_graves"]:
            detalle.append(f"- Errores graves: {h['comentario_graves']}")
        detalle += [f"- {k}: {t}" for k, t in h["textos"].items()]
        if detalle:
            partes += ["", f"Comentarios de `{h['fichero']}`:", *detalle]
    return "\n".join(partes)


def escribir_valoracion_en_ficha(carpeta, v, metricas):
    ficha = carpeta / "ficha.md"
    if not ficha.exists():
        return
    txt = ficha.read_text(encoding="utf-8")
    bloque = f"{MARCA_VAL_INI}\n{bloque_valoracion(v, metricas)}\n{MARCA_VAL_FIN}"
    if MARCA_VAL_INI in txt:
        txt = re.sub(re.escape(MARCA_VAL_INI) + ".*?" + re.escape(MARCA_VAL_FIN), lambda _: bloque, txt, flags=re.S)
    else:
        txt = txt.replace("## Pendiente de revisión experta", f"{SECCION_VALORACION}\n\n{bloque}\n\n"
                          "## Pendiente de revisión experta", 1)
    ficha.write_text(txt, encoding="utf-8")
    actualizar_ficha(carpeta, metricas)


def cmd_valoracion(args):
    carpeta = buscar_ejecucion(args.ejecucion)
    metricas = leer_json(carpeta / "metricas.json")
    if metricas is None:
        error(f"falta {carpeta.name}/metricas.json: ejecuta antes `medir`")
    rutas = sorted(carpeta.glob("valoracion-coordinacion*.xlsx"))
    if not rutas:
        error(f"no hay valoracion-coordinacion*.xlsx en {carpeta.name}")
    v = consolidar_valoracion([leer_hoja_valoracion(r) for r in rutas], metricas)
    (carpeta / "valoracion.json").write_text(json.dumps(v, ensure_ascii=False, indent=2), encoding="utf-8")
    escribir_valoracion_en_ficha(carpeta, v, metricas)
    actualizar_historial()
    print(bloque_valoracion(v, metricas))
    print(f"\nActualizados {carpeta.name}/valoracion.json, ficha.md y evals/historial.md")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="orden", required=True)
    e = sub.add_parser("enlaces", help="lista los enlaces a documentos y páginas de calidad de una página web")
    e.add_argument("url")
    e.add_argument("--todos", action="store_true", help="muestra todos los enlaces, sin filtrar")
    e.set_defaults(func=cmd_enlaces)
    c = sub.add_parser("caso", help="da de alta un caso y descarga sus documentos desde un manifiesto JSON")
    c.add_argument("manifiesto", help="p. ej. evals/casos/<id>/fuentes.json")
    c.add_argument("--sin-descarga", action="store_true",
                   help="los documentos ya están copiados a mano en entrada/ y referencia/")
    c.set_defaults(func=cmd_caso)
    n = sub.add_parser("nueva", help="crea la carpeta y la ficha de una ejecución")
    n.add_argument("id")
    n.add_argument("--modelo-a", default=MODELO_A)
    n.add_argument("--modelo-b", default=MODELO_B)
    n.set_defaults(func=cmd_nueva)
    m = sub.add_parser("medir", help="calcula métricas, consolida cotejos y actualiza el historial")
    m.add_argument("ejecucion", help="prefijo de la ejecución, p. ej. E04")
    m.set_defaults(func=cmd_medir)
    v = sub.add_parser("valoracion", help="incorpora las hojas de valoración devueltas por la coordinación")
    v.add_argument("ejecucion", help="prefijo de la ejecución, p. ej. E04")
    v.set_defaults(func=cmd_valoracion)
    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
